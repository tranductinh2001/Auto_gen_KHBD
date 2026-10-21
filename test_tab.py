import docx
from docx.oxml import parse_xml
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.left_margin = Inches(0.79)   # 2.0 cm
sec.right_margin = Inches(0.59)  # 1.5 cm

# Test Tab leader dots
for i in range(5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    pPr = p._p.get_or_add_pPr()
    tabs_xml = '<w:tabs xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:tab w:val="right" w:leader="dot" w:pos="9920"/></w:tabs>'
    pPr.append(parse_xml(tabs_xml))
    r = p.add_run("\t")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

doc.save("test_tab.docx")
print("SUCCESS generating test_tab.docx!")
