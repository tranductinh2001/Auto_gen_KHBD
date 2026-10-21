import os
import csv
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Tuple


class GraphDataExporter:
    """Xuất bản dữ liệu đồ thị tri thức sang cả Excel đa Sheet và các tệp CSV chuẩn hóa cho Neo4j."""

    def __init__(self, output_dir: str = "output"):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def _setup_header(self, ws, headers: list, fill_color: str = "1F4E79", font_color: str = "FFFFFF"):
        header_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        header_font = Font(bold=True, color=font_color, size=11)
        thin_border = Border(
            left=Side(style='thin', color='D3D3D3'),
            right=Side(style='thin', color='D3D3D3'),
            top=Side(style='thin', color='D3D3D3'),
            bottom=Side(style='thin', color='A0A0A0')
        )
        ws.row_dimensions[1].height = 28
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
            col_letter = openpyxl.utils.get_column_letter(col_num)
            ws.column_dimensions[col_letter].width = max(len(str(header)) + 5, 16)

    def export(
        self,
        extracted_bundles: List[dict],
        logs: List[dict],
        base_name: str = "tinhoc_gdpt_2018"
    ) -> Tuple[str, str, str]:
        """
        Xuất toàn bộ dữ liệu:
        1. File Excel (.xlsx) đa Sheet của phiên cào này.
        2. Thư mục CSV riêng của phiên cào này: output/sessions/{base_name}_{timestamp}/ (không bị ghi đè).
        3. Thư mục Master Neo4j: output/neo4j_import/ (tự động gộp tích lũy và loại trùng lặp).
        """
        all_flat = []
        all_lessons = []
        all_requirements = []
        all_knowledges = []
        all_visual_assets = []
        all_skills = []
        all_questions = []
        all_answers = []

        seen_lessons = set()
        seen_requirements = set()
        seen_knowledges = set()
        seen_visual_assets = set()
        seen_skills = set()
        seen_questions = set()

        for b in extracted_bundles:
            all_flat.extend(b.get("flat_rows", []))
            norm = b.get("normalized", {})

            for item in norm.get("lessons", []):
                k = (item.get("grade_level"), item.get("lesson_id"))
                if k not in seen_lessons:
                    seen_lessons.add(k)
                    all_lessons.append(item)

            for item in norm.get("requirements", []):
                k = (item.get("grade_level"), item.get("lesson_id"), item.get("requirement_code"))
                if k not in seen_requirements:
                    seen_requirements.add(k)
                    all_requirements.append(item)

            for item in norm.get("knowledges", []):
                k = (item.get("grade_level"), item.get("lesson_id"), item.get("knowledge_name"))
                if k not in seen_knowledges:
                    seen_knowledges.add(k)
                    all_knowledges.append(item)

            for item in norm.get("visual_assets", []):
                k = (item.get("grade_level"), item.get("lesson_id"), item.get("visual_asset_code"))
                if k not in seen_visual_assets:
                    seen_visual_assets.add(k)
                    all_visual_assets.append(item)

            for item in norm.get("skills", []):
                k = (item.get("grade_level"), item.get("lesson_id"), item.get("skill_name"))
                if k not in seen_skills:
                    seen_skills.add(k)
                    all_skills.append(item)

            for item in norm.get("questions", []):
                k = (item.get("grade_level"), item.get("lesson_id"), item.get("question_id"))
                if k not in seen_questions:
                    seen_questions.add(k)
                    all_questions.append(item)

            for item in norm.get("answers", []):
                all_answers.append(item)

        # ── 1. Tạo File Excel đa Sheet cho phiên này ──
        timestamp_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        excel_filename = f"{base_name}_{timestamp_str}.xlsx"
        session_excel_path = os.path.join(self.output_dir, excel_filename)

        wb = openpyxl.Workbook()
        self._write_flat_sheet(wb, all_flat)
        self._write_visual_assets_sheet(wb, all_visual_assets)
        self._write_lessons_sheet(wb, all_lessons)
        self._write_requirements_sheet(wb, all_requirements)
        self._write_knowledges_sheet(wb, all_knowledges)
        self._write_skills_sheet(wb, all_skills)
        self._write_questions_sheet(wb, all_questions, all_answers)
        self._write_logs_sheet(wb, logs)
        wb.save(session_excel_path)

        # ── 2. Xuất CSV riêng cho phiên này (output/sessions/...) ──
        session_dir = os.path.join(self.output_dir, "sessions", f"{base_name}_{timestamp_str}")
        os.makedirs(session_dir, exist_ok=True)
        self._write_csv_bundle(
            target_dir=session_dir,
            flat=all_flat,
            lessons=all_lessons,
            requirements=all_requirements,
            knowledges=all_knowledges,
            visual_assets=all_visual_assets,
            skills=all_skills,
            questions=all_questions,
            answers=all_answers
        )

        # ── 3. Tích lũy và Gộp khử trùng lặp vào Kho Master Neo4j ──
        master_dir = os.path.join(self.output_dir, "neo4j_import")
        os.makedirs(master_dir, exist_ok=True)
        self._accumulate_to_master(
            master_dir=master_dir,
            new_flat=all_flat,
            new_lessons=all_lessons,
            new_requirements=all_requirements,
            new_knowledges=all_knowledges,
            new_visual_assets=all_visual_assets,
            new_skills=all_skills,
            new_questions=all_questions,
            new_answers=all_answers
        )

        return session_excel_path, session_dir, master_dir

    def _write_csv_bundle(self, target_dir: str, flat, lessons, requirements, knowledges, visual_assets, skills, questions, answers):
        def save_csv(filename, fields, data):
            filepath = os.path.join(target_dir, filename)
            with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.DictWriter(f, fieldnames=fields, extrasaction='ignore')
                writer.writeheader()
                for row in data:
                    writer.writerow(row)

        flat_cols = [
            "grade_level", "chapter_title", "lesson_id", "lesson_title",
            "requirement_code", "requirement_content", "knowledge_name",
            "knowledge_content", "skill_name", "skill_description",
            "visual_asset_code", "caption", "detailed_description",
            "question_id", "content", "difficulty", "option_key",
            "option_content", "is_correct", "page_num", "source_file"
        ]
        save_csv("neo4j_flat_dataset.csv", flat_cols, flat)
        save_csv("neo4j_lessons.csv", ["grade_level", "chapter_title", "lesson_id", "lesson_title", "page_num", "source_file"], lessons)
        save_csv("neo4j_requirements.csv", ["grade_level", "chapter_title", "lesson_id", "lesson_title", "requirement_code", "requirement_content", "page_num"], requirements)
        save_csv("neo4j_knowledges.csv", ["grade_level", "chapter_title", "lesson_id", "lesson_title", "knowledge_name", "knowledge_content", "page_num"], knowledges)
        save_csv("neo4j_visual_assets.csv", ["grade_level", "chapter_title", "lesson_id", "lesson_title", "visual_asset_code", "caption", "detailed_description", "page_num"], visual_assets)
        save_csv("neo4j_skills.csv", ["grade_level", "chapter_title", "lesson_id", "lesson_title", "skill_name", "skill_description", "page_num"], skills)
        save_csv("neo4j_questions.csv", ["grade_level", "chapter_title", "lesson_id", "lesson_title", "question_id", "content", "difficulty", "page_num"], questions)
        save_csv("neo4j_answer_options.csv", ["grade_level", "lesson_id", "question_id", "option_key", "content", "is_correct", "page_num"], answers)

        # Sinh Cypher script nạp dữ liệu
        cypher_file = os.path.join(target_dir, "neo4j_import_script.cypher")
        cypher_content = """// ==============================================================================
// NEO4J KNOWLEDGE GRAPH IMPORT SCRIPT: TIN HỌC PHỔ THÔNG GDPT 2018
// Đề tài: Quản lý tri thức, phân tích chương trình & hỗ trợ ra đề thi tự động
// ==============================================================================

// 1. TẠO RÀNG BUỘC KHÓA DUY NHẤT & CHỈ MỤC
CREATE CONSTRAINT constraint_grade IF NOT EXISTS FOR (g:Grade) REQUIRE g.grade_level IS UNIQUE;
CREATE CONSTRAINT constraint_chapter IF NOT EXISTS FOR (c:Chapter) REQUIRE c.id IS UNIQUE;
CREATE CONSTRAINT constraint_lesson IF NOT EXISTS FOR (l:Lesson) REQUIRE l.id IS UNIQUE;
CREATE CONSTRAINT constraint_req IF NOT EXISTS FOR (r:Requirement) REQUIRE r.id IS UNIQUE;
CREATE CONSTRAINT constraint_know IF NOT EXISTS FOR (k:Knowledge) REQUIRE k.id IS UNIQUE;
CREATE CONSTRAINT constraint_visual IF NOT EXISTS FOR (v:VisualAsset) REQUIRE v.id IS UNIQUE;
CREATE CONSTRAINT constraint_skill IF NOT EXISTS FOR (s:Skill) REQUIRE s.id IS UNIQUE;
CREATE CONSTRAINT constraint_question IF NOT EXISTS FOR (q:Question) REQUIRE q.id IS UNIQUE;

// 2. NẠP CẤU TRÚC PHÂN CẤP SƯ PHẠM (GRADE -> CHAPTER -> LESSON)
LOAD CSV WITH HEADERS FROM 'file:///neo4j_lessons.csv' AS row
WITH row WHERE row.lesson_id IS NOT NULL AND trim(row.lesson_id) <> ''
MERGE (g:Grade {grade_level: coalesce(row.grade_level, 'Lớp 9')})
MERGE (c:Chapter {id: coalesce(row.grade_level, 'Lớp 9') + '_' + coalesce(row.chapter_title, 'Chủ đề chung')})
ON CREATE SET c.title = coalesce(row.chapter_title, 'Chủ đề chung')
MERGE (g)-[:HAS_CHAPTER]->(c)
MERGE (l:Lesson {id: coalesce(row.grade_level, 'Lớp 9') + '_' + row.lesson_id})
ON CREATE SET l.lesson_id = row.lesson_id,
              l.title = row.lesson_title,
              l.page = row.page_num,
              l.source_file = row.source_file
MERGE (c)-[:HAS_LESSON]->(l);

// 3. NẠP YÊU CẦU CẦN ĐẠT (REQUIREMENTS)
LOAD CSV WITH HEADERS FROM 'file:///neo4j_requirements.csv' AS row
WITH row WHERE row.requirement_code IS NOT NULL AND trim(row.requirement_code) <> ''
MATCH (l:Lesson {id: coalesce(row.grade_level, 'Lớp 9') + '_' + row.lesson_id})
MERGE (r:Requirement {id: coalesce(row.grade_level, 'Lớp 9') + '_' + row.lesson_id + '_' + row.requirement_code})
ON CREATE SET r.code = row.requirement_code,
              r.content = row.requirement_content
MERGE (l)-[:HAS_REQUIREMENT]->(r);

// 4. NẠP NỘI DUNG KIẾN THỨC CỐT LÕI (KNOWLEDGE)
LOAD CSV WITH HEADERS FROM 'file:///neo4j_knowledges.csv' AS row
WITH row WHERE row.knowledge_name IS NOT NULL AND trim(row.knowledge_name) <> ''
MATCH (l:Lesson {id: coalesce(row.grade_level, 'Lớp 9') + '_' + row.lesson_id})
MERGE (k:Knowledge {id: coalesce(row.grade_level, 'Lớp 9') + '_' + row.lesson_id + '_' + row.knowledge_name})
ON CREATE SET k.name = row.knowledge_name,
              k.content = row.knowledge_content
MERGE (l)-[:TEACHES_KNOWLEDGE]->(k);

// 5. NẠP HÌNH ẢNH, SƠ ĐỒ VÀ GIAO DIỆN (MULTIMODAL VISUAL ASSETS)
LOAD CSV WITH HEADERS FROM 'file:///neo4j_visual_assets.csv' AS row
WITH row WHERE row.visual_asset_code IS NOT NULL AND trim(row.visual_asset_code) <> ''
MATCH (l:Lesson {id: coalesce(row.grade_level, 'Lớp 9') + '_' + row.lesson_id})
MERGE (v:VisualAsset {id: coalesce(row.grade_level, 'Lớp 9') + '_' + row.lesson_id + '_' + row.visual_asset_code})
ON CREATE SET v.code = row.visual_asset_code,
              v.caption = row.caption,
              v.detailed_description = row.detailed_description,
              v.page = row.page_num
MERGE (l)-[:HAS_VISUAL_ASSET]->(v);

// 6. NẠP KỸ NĂNG THỰC HÀNH (SKILLS)
LOAD CSV WITH HEADERS FROM 'file:///neo4j_skills.csv' AS row
WITH row WHERE row.skill_name IS NOT NULL AND trim(row.skill_name) <> ''
MATCH (l:Lesson {id: coalesce(row.grade_level, 'Lớp 9') + '_' + row.lesson_id})
MERGE (s:Skill {id: coalesce(row.grade_level, 'Lớp 9') + '_' + row.lesson_id + '_' + row.skill_name})
ON CREATE SET s.name = row.skill_name,
              s.description = row.skill_description
MERGE (l)-[:DEVELOPS_SKILL]->(s);

// 7. NẠP CÂU HỎI ĐÁNH GIÁ (QUESTIONS)
LOAD CSV WITH HEADERS FROM 'file:///neo4j_questions.csv' AS row
WITH row WHERE row.question_id IS NOT NULL AND trim(row.question_id) <> ''
MATCH (l:Lesson {id: coalesce(row.grade_level, 'Lớp 9') + '_' + row.lesson_id})
MERGE (q:Question {id: coalesce(row.grade_level, 'Lớp 9') + '_' + row.lesson_id + '_' + row.question_id})
ON CREATE SET q.question_id = row.question_id,
              q.content = row.content,
              q.difficulty = row.difficulty
MERGE (l)-[:EVALUATED_BY]->(q);

// 8. NẠP CÁC PHƯƠNG ÁN LỰA CHỌN & ĐÁP ÁN ĐÚNG (ANSWER OPTIONS)
LOAD CSV WITH HEADERS FROM 'file:///neo4j_answer_options.csv' AS row
WITH row WHERE row.question_id IS NOT NULL AND trim(row.question_id) <> ''
MATCH (q:Question {id: coalesce(row.grade_level, 'Lớp 9') + '_' + row.lesson_id + '_' + row.question_id})
CREATE (a:AnswerOption {
    option_key: coalesce(row.option_key, ''),
    content: row.content,
    is_correct: (row.is_correct = 'True' OR row.is_correct = 'true')
})
MERGE (q)-[:HAS_OPTION]->(a);
"""
        with open(cypher_file, "w", encoding="utf-8") as f:
            f.write(cypher_content)

    def _accumulate_to_master(self, master_dir: str, new_flat, new_lessons, new_requirements, new_knowledges, new_visual_assets, new_skills, new_questions, new_answers):
        """Hợp nhất tích lũy dữ liệu mới vào kho tổng hợp Master Neo4j, tự động loại bỏ trùng lặp."""

        def merge_records(filename, new_items, key_fn):
            path = os.path.join(master_dir, filename)
            existing = []
            if os.path.exists(path):
                try:
                    with open(path, "r", encoding="utf-8-sig") as f:
                        reader = csv.DictReader(f)
                        for r in reader:
                            existing.append(dict(r))
                except Exception:
                    pass

            merged = []
            seen = set()
            for r in existing:
                k = key_fn(r)
                if k not in seen:
                    seen.add(k)
                    merged.append(r)

            for r in new_items:
                k = key_fn(r)
                if k not in seen:
                    seen.add(k)
                    merged.append(r)

            return merged

        m_lessons = merge_records("neo4j_lessons.csv", new_lessons, lambda r: (r.get("grade_level", "").strip(), r.get("lesson_id", "").strip()))
        m_requirements = merge_records("neo4j_requirements.csv", new_requirements, lambda r: (r.get("grade_level", "").strip(), r.get("lesson_id", "").strip(), r.get("requirement_code", "").strip()))
        m_knowledges = merge_records("neo4j_knowledges.csv", new_knowledges, lambda r: (r.get("grade_level", "").strip(), r.get("lesson_id", "").strip(), r.get("knowledge_name", "").strip()))
        m_visual_assets = merge_records("neo4j_visual_assets.csv", new_visual_assets, lambda r: (r.get("grade_level", "").strip(), r.get("lesson_id", "").strip(), r.get("visual_asset_code", "").strip()))
        m_skills = merge_records("neo4j_skills.csv", new_skills, lambda r: (r.get("grade_level", "").strip(), r.get("lesson_id", "").strip(), r.get("skill_name", "").strip()))
        m_questions = merge_records("neo4j_questions.csv", new_questions, lambda r: (r.get("grade_level", "").strip(), r.get("lesson_id", "").strip(), r.get("question_id", "").strip()))
        m_answers = merge_records("neo4j_answer_options.csv", new_answers, lambda r: (r.get("grade_level", "").strip(), r.get("lesson_id", "").strip(), r.get("question_id", "").strip(), r.get("option_key", "").strip()))
        m_flat = merge_records("neo4j_flat_dataset.csv", new_flat, lambda r: (
            r.get("grade_level", "").strip(), r.get("lesson_id", "").strip(), r.get("requirement_code", "").strip(),
            r.get("knowledge_name", "").strip(), r.get("visual_asset_code", "").strip(), r.get("skill_name", "").strip(),
            r.get("question_id", "").strip(), r.get("option_key", "").strip()
        ))

        self._write_csv_bundle(master_dir, m_flat, m_lessons, m_requirements, m_knowledges, m_visual_assets, m_skills, m_questions, m_answers)

        # Xuất file Excel Master tổng hợp
        master_excel = os.path.join(self.output_dir, "sgk_tinhoc_master_knowledge_graph.xlsx")
        wb_master = openpyxl.Workbook()
        self._write_flat_sheet(wb_master, m_flat)
        self._write_visual_assets_sheet(wb_master, m_visual_assets)
        self._write_lessons_sheet(wb_master, m_lessons)
        self._write_requirements_sheet(wb_master, m_requirements)
        self._write_knowledges_sheet(wb_master, m_knowledges)
        self._write_skills_sheet(wb_master, m_skills)
        self._write_questions_sheet(wb_master, m_questions, m_answers)
        wb_master.save(master_excel)

    # ─── Sheet Helpers ──────────────────────────────────────────

    def _write_flat_sheet(self, wb, flat: list):
        ws = wb.active
        ws.title = "DuLieu_TongHop"
        headers = [
            "grade_level", "chapter_title", "lesson_id", "lesson_title",
            "requirement_code", "requirement_content", "knowledge_name",
            "knowledge_content", "skill_name", "skill_description",
            "visual_asset_code", "caption", "detailed_description",
            "question_id", "content", "difficulty", "option_key",
            "option_content", "is_correct", "page_num", "source_file"
        ]
        self._setup_header(ws, headers, fill_color="1F4E79", font_color="FFFFFF")
        for row_idx, r in enumerate(flat, start=2):
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=str(r.get(h, "")))
                cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if h in ["grade_level", "lesson_id", "requirement_code", "visual_asset_code", "question_id", "difficulty", "option_key", "is_correct", "page_num"] else "left")

    def _write_visual_assets_sheet(self, wb, visual_assets: list):
        ws = wb.create_sheet(title="HinhAnh_SoDo_GiaoDien")
        headers = ["grade_level", "chapter_title", "lesson_id", "lesson_title", "visual_asset_code", "caption", "detailed_description", "page_num"]
        self._setup_header(ws, headers, fill_color="7030A0", font_color="FFFFFF")
        ws.column_dimensions['E'].width = 16
        ws.column_dimensions['F'].width = 32
        ws.column_dimensions['G'].width = 65
        for row_idx, v in enumerate(visual_assets, start=2):
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=str(v.get(h, "")))
                cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if h in ["grade_level", "lesson_id", "visual_asset_code", "page_num"] else "left")

    def _write_lessons_sheet(self, wb, lessons: list):
        ws = wb.create_sheet(title="BaiHoc")
        headers = ["grade_level", "chapter_title", "lesson_id", "lesson_title", "page_num", "source_file"]
        self._setup_header(ws, headers, fill_color="2F5597", font_color="FFFFFF")
        for row_idx, l in enumerate(lessons, start=2):
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=str(l.get(h, "")))
                cell.alignment = Alignment(vertical="center", horizontal="center" if h in ["grade_level", "lesson_id", "page_num"] else "left")

    def _write_requirements_sheet(self, wb, requirements: list):
        ws = wb.create_sheet(title="YeuCauCanDat")
        headers = ["grade_level", "chapter_title", "lesson_id", "lesson_title", "requirement_code", "requirement_content", "page_num"]
        self._setup_header(ws, headers, fill_color="31869B", font_color="FFFFFF")
        for row_idx, r in enumerate(requirements, start=2):
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=str(r.get(h, "")))
                cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if h in ["grade_level", "lesson_id", "requirement_code", "page_num"] else "left")

    def _write_knowledges_sheet(self, wb, knowledges: list):
        ws = wb.create_sheet(title="KienThuc")
        headers = ["grade_level", "chapter_title", "lesson_id", "lesson_title", "knowledge_name", "knowledge_content", "page_num"]
        self._setup_header(ws, headers, fill_color="375623", font_color="FFFFFF")
        for row_idx, k in enumerate(knowledges, start=2):
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=str(k.get(h, "")))
                cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if h in ["grade_level", "lesson_id", "page_num"] else "left")

    def _write_skills_sheet(self, wb, skills: list):
        ws = wb.create_sheet(title="KyNang")
        headers = ["grade_level", "chapter_title", "lesson_id", "lesson_title", "skill_name", "skill_description", "page_num"]
        self._setup_header(ws, headers, fill_color="C65911", font_color="FFFFFF")
        for row_idx, s in enumerate(skills, start=2):
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=str(s.get(h, "")))
                cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if h in ["grade_level", "lesson_id", "page_num"] else "left")

    def _write_questions_sheet(self, wb, questions: list, answers: list):
        ws = wb.create_sheet(title="NganHang_CauHoi_DeThi")
        headers = ["grade_level", "lesson_id", "lesson_title", "question_id", "content", "difficulty", "option_key", "option_content", "is_correct", "page_num"]
        self._setup_header(ws, headers, fill_color="833C0C", font_color="FFFFFF")
        q_map = {(q.get("grade_level"), q.get("lesson_id"), q.get("question_id")): q for q in questions}

        row_idx = 2
        for a in answers:
            key = (a.get("grade_level"), a.get("lesson_id"), a.get("question_id"))
            q_info = q_map.get(key, {})
            row_data = [
                a.get("grade_level", ""),
                a.get("lesson_id", ""),
                q_info.get("lesson_title", ""),
                a.get("question_id", ""),
                q_info.get("content", ""),
                q_info.get("difficulty", ""),
                a.get("option_key", ""),
                a.get("content", ""),
                str(a.get("is_correct", "")),
                a.get("page_num", "")
            ]
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=str(val))
                cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if col_idx in [1, 2, 4, 6, 7, 9, 10] else "left")
            row_idx += 1

    def _write_logs_sheet(self, wb, logs: list):
        ws = wb.create_sheet(title="NhatKy_TienTrinh")
        headers = ["source_file", "page_num", "status", "rows_extracted", "duration_sec", "error"]
        self._setup_header(ws, headers, fill_color="595959", font_color="FFFFFF")
        for row_idx, l in enumerate(logs, start=2):
            for col_idx, h in enumerate(headers, start=1):
                ws.cell(row=row_idx, column=col_idx, value=str(l.get(h, "")))
