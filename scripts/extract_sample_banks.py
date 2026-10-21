import os
import sys
import re
import json
import docx

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

def parse_full_exam_doc(filepath, grade_key):
    doc = docx.Document(filepath)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    
    # 1. Parse tables for answers
    ans_p1 = {}
    ans_p2 = {}
    ans_p3 = []
    
    for t in doc.tables:
        if len(t.rows) >= 2:
            r0 = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells]
            if 'CÂU' in r0 or 'Câu' in r0:
                if 'ĐÁP ÁN' in [c.text.strip().replace('\n', ' ') for c in t.rows[1].cells]:
                    cols = len(r0)
                    r1 = [c.text.strip().replace('\n', ' ') for c in t.rows[1].cells]
                    for idx in range(1, cols):
                        q_num = r0[idx].strip()
                        q_ans = r1[idx].strip()
                        if q_num.isdigit() and q_ans:
                            ans_p1[int(q_num)] = q_ans
            if 'Lệnh hỏi' in r0 and 'Đáp án (Đ/S)' in r0:
                for r in t.rows[1:]:
                    cells = [c.text.strip() for c in r.cells]
                    if len(cells) >= 3 and cells[0].isdigit() and cells[1] and cells[2]:
                        q_num = int(cells[0])
                        sub = cells[1].upper()
                        val = cells[2]
                        if q_num not in ans_p2: ans_p2[q_num] = {}
                        ans_p2[q_num][sub] = val
                    if len(cells) >= 6 and cells[3].isdigit() and cells[4] and cells[5]:
                        q_num = int(cells[3])
                        sub = cells[4].upper()
                        val = cells[5]
                        if q_num not in ans_p2: ans_p2[q_num] = {}
                        ans_p2[q_num][sub] = val
            if 'Nội dung' in r0 or 'Đáp án' in r0:
                for r in t.rows[1:]:
                    cells = [c.text.strip() for c in r.cells]
                    if len(cells) >= 3:
                        ans_p3.append({'cau': cells[0], 'noi_dung': cells[1], 'diem': cells[2]})
                        
    # 2. Extract Questions
    data = {'grade': grade_key, 'p1_mcq': [], 'p2_tf': [], 'p3_essay': [], 'answers': {'p1': ans_p1, 'p2': ans_p2, 'p3': ans_p3}}
    
    current_mode = None
    curr_q = None
    
    for p in paras:
        if any(k in p.upper() for k in ['I. KHOANH TRÒN', 'PHẦN A: CÂU TRẮC NGHIỆM NHIỀU PHƯƠNG ÁN', 'I. DẠNG NHIỀU LỰA CHỌN', 'I. CÂU HỎI TRẮC NGHIỆM']):
            current_mode = 'P1'
            continue
        elif any(k in p.upper() for k in ['II. VIẾT CHỮ Đ', 'PHẦN B: CÂU TRẮC NGHIỆM ĐÚNG/SAI', 'II. DẠNG ĐÚNG/SAI']):
            current_mode = 'P2'
            continue
        elif any(k in p.upper() for k in ['PHẦN C:  TỰ LUẬN', 'B. PHẦN TỰ LUẬN', 'PHẦN TỰ LUẬN', 'B. PHẦN TỰ LUẬN:']):
            current_mode = 'P3'
            continue
        elif any(k in p.upper() for k in ['BÀI LÀM', 'HƯỚNG DẪN CHẤM BIỂU ĐIỂM', '------HẾT------']):
            if current_mode == 'P3':
                current_mode = None
            continue
            
        if current_mode == 'P1':
            m = re.match(r'^Câu\s+(\d+)[\.\:](.*)', p)
            if m:
                if curr_q and 'options' in curr_q:
                    data['p1_mcq'].append(curr_q)
                curr_q = {'num': int(m.group(1)), 'text': m.group(2).strip(), 'options': []}
            elif curr_q and 'options' in curr_q:
                curr_q['options'].append(p)
        elif current_mode == 'P2':
            m = re.match(r'^Câu\s+(\d+)[\.\:](.*)', p)
            if m:
                if curr_q and curr_q.get('type') == 'TF':
                    data['p2_tf'].append(curr_q)
                curr_q = {'type': 'TF', 'num': int(m.group(1)), 'text': m.group(2).strip(), 'items': []}
            elif curr_q and curr_q.get('type') == 'TF':
                curr_q['items'].append(p)
        elif current_mode == 'P3':
            m = re.match(r'^Câu\s+(\d+)[\.\:]?(.*)', p)
            if m:
                if curr_q and curr_q.get('type') == 'ESSAY':
                    data['p3_essay'].append(curr_q)
                curr_q = {'type': 'ESSAY', 'num': int(m.group(1)), 'text': m.group(2).strip()}
            elif curr_q and curr_q.get('type') == 'ESSAY':
                curr_q['text'] += ' ' + p
                
    if curr_q:
        if curr_q.get('type') == 'ESSAY':
            data['p3_essay'].append(curr_q)
        elif curr_q.get('type') == 'TF':
            data['p2_tf'].append(curr_q)
        elif 'options' in curr_q:
            data['p1_mcq'].append(curr_q)
            
    print(f"Extracted {grade_key}: {len(data['p1_mcq'])} MCQs, {len(data['p2_tf'])} T/Fs, {len(data['p3_essay'])} Essays")
    return data

if __name__ == '__main__':
    bank = {}
    bank['Lớp 9'] = parse_full_exam_doc('để thi mẫu/Le Thi Phuong Thao_DKT GHK 1_25-26_ TIN HỌC 9.docx', 'Lớp 9')
    bank['Lớp 7'] = parse_full_exam_doc('để thi mẫu/KTGK1 (25- 26) - TIN 7  (02) - PHƯƠNG THẢO (1).docx', 'Lớp 7')
    bank['Lớp 6'] = parse_full_exam_doc('để thi mẫu/Le Thi Phuong Thao_DKT GHK 1_25-26_ TIN HỌC 6.docx', 'Lớp 6')
    
    out_path = 'data_input/extracted_sample_banks.json'
    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(bank, f, ensure_ascii=False, indent=2)
    print(f"Saved complete banks to {out_path}")
