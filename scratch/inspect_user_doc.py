import os, sys, glob, shutil
import docx
sys.stdout.reconfigure(encoding='utf-8')

dl_dir = 'C:/Users/Admin/Downloads'
files = [os.path.join(dl_dir, f) for f in os.listdir(dl_dir) if f.startswith('Ho_So_De_Thi_Chuan') and f.endswith('.docx')]
files.sort(key=os.path.getmtime, reverse=True)
if not files:
    print("No files found")
    sys.exit(0)

target = files[0]
print("Found latest:", target)
copy_to = os.path.abspath('scratch/user_downloaded_exam.docx')
shutil.copyfile(target, copy_to)
print("Copied to:", copy_to)

doc = docx.Document(copy_to)
print(f"Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}, Total sections: {len(doc.sections)}")

# Let's inspect sections and text
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith("I.") or txt.startswith("II.") or txt.startswith("Câu ") or txt.startswith("BÀI LÀM"):
        # Check runs font size
        fs_list = [r.font.size.pt if r.font.size else 'None' for r in p.runs]
        print(f"P{i:3d}: {txt[:60]}... | font_sizes={set(fs_list)}")
