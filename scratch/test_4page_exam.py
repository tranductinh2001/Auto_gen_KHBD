import docx
import win32com.client
import os
import sys

sys.path.insert(0, os.getcwd())
sys.stdout.reconfigure(encoding='utf-8')

from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from core.exam_builder import ExamMatrixEngine, ExamDocxExporter

engine = ExamMatrixEngine(
    grade_level='Lớp 9',
    exam_title='KIỂM TRA GIỮA HỌC KỲ I',
    school_name='TRƯỜNG TH&THCS ĐĂK TRÔI',
    governing_body='UBND XÃ KON CHIÊNG',
    school_year='2026 - 2027',
    teacher_name='Lê Thị Phương Thảo',
    head_name='Trần Thị Tuyết Ngân'
)
dossier = engine.generate_exam_dossier()

# We will test rendering TH01 directly into a fresh document to check page count
doc = docx.Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.PORTRAIT
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(0.59)
sec.bottom_margin = Inches(0.59)
sec.left_margin = Inches(0.59)
sec.right_margin = Inches(0.59)

meta = dossier["meta"]
exam_data = dossier["th01"]
code = exam_data["code"]

# ==============================================================================
# TRANG 1: MẶT TRƯỚC TỜ 1 (TRẮC NGHIỆM CÂU 1 ĐẾN CÂU 8)
# ==============================================================================
ExamDocxExporter._render_header_box(doc, meta, code, is_essay=False)
ExamDocxExporter._render_cut_line(doc)
ExamDocxExporter._render_score_box(doc, code, is_essay=False)

p_p1 = doc.add_paragraph()
p_p1.paragraph_format.space_before = Pt(3)
p_p1.paragraph_format.space_after = Pt(2)
p_p1.paragraph_format.keep_with_next = True
r_p1 = p_p1.add_run("I. Khoanh tròn vào chữ cái đứng trước câu trả lời đúng (3,0 điểm)")
r_p1.font.name = "Times New Roman"
r_p1.font.size = Pt(12.5)
r_p1.bold = True

PAGE1_COUNT = 8

def render_mcq_compact(q, idx):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(2.5)
    p_q.paragraph_format.space_after = Pt(1.0)
    p_q.paragraph_format.line_spacing = 1.10
    p_q.paragraph_format.keep_with_next = True

    rq_num = p_q.add_run(f"Câu {idx}. ")
    rq_num.font.name = "Times New Roman"
    rq_num.font.size = Pt(11.5)
    rq_num.bold = True

    rq_cnt = p_q.add_run(q["content"])
    rq_cnt.font.name = "Times New Roman"
    rq_cnt.font.size = Pt(11.5)

    # Render options: threshold 54 chars for 2 per line
    opts = q["options"]
    opt_keys = ["A", "B", "C", "D"]
    max_opt_len = max(len(str(opts.get(k, ""))) for k in opt_keys)

    if max_opt_len <= 14:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.20)
        p.paragraph_format.space_before = Pt(0.5)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.10
        p.paragraph_format.keep_with_next = False
        p.paragraph_format.tab_stops.add_tab_stop(Inches(1.85), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.60), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(5.35), WD_TAB_ALIGNMENT.LEFT)
        for i, k in enumerate(opt_keys):
            if i > 0: p.add_run("\t")
            p.add_run(f"{k}. ").bold = True
            p.add_run(str(opts.get(k, "")))
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11.5)
    elif max_opt_len <= 54:
        for p_idx, (k1, k2) in enumerate([("A", "B"), ("C", "D")]):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.20)
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.10
            p.paragraph_format.keep_with_next = (p_idx == 0)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(3.60), WD_TAB_ALIGNMENT.LEFT)

            r1 = p.add_run(f"{k1}. ")
            r1.bold = True
            p.add_run(str(opts.get(k1, "")))
            p.add_run("\t")
            r2 = p.add_run(f"{k2}. ")
            r2.bold = True
            p.add_run(str(opts.get(k2, "")))
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11.5)
    else:
        for i, k in enumerate(opt_keys):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.20)
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(1.0)
            p.paragraph_format.line_spacing = 1.10
            p.paragraph_format.keep_with_next = (i < 3)
            r = p.add_run(f"{k}. ")
            r.bold = True
            p.add_run(str(opts.get(k, "")))
            for run_i in p.runs:
                run_i.font.name = "Times New Roman"
                run_i.font.size = Pt(11.5)

# Render Page 1 MCQs
for idx, q in enumerate(exam_data["mcqs"][:PAGE1_COUNT], 1):
    render_mcq_compact(q, idx)

# ==============================================================================
# TRANG 2: MẶT SAU TỜ 1 (TRẮC NGHIỆM TIẾP THEO + PHẦN II ĐÚNG/SAI)
# ==============================================================================
doc.add_page_break()
ExamDocxExporter._render_x_cut_box(doc)

for idx, q in enumerate(exam_data["mcqs"][PAGE1_COUNT:], PAGE1_COUNT + 1):
    render_mcq_compact(q, idx)

# Phần II
p_p2 = doc.add_paragraph()
p_p2.paragraph_format.space_before = Pt(2.5)
p_p2.paragraph_format.space_after = Pt(1.5)
p_p2.paragraph_format.keep_with_next = True
r_p2 = p_p2.add_run("II. Viết chữ Đ(đúng) hoặc S(sai) vào các câu trả lời tương ứng trong các câu sau: (4,0 điểm)")
r_p2.font.name = "Times New Roman"
r_p2.font.size = Pt(12)
r_p2.bold = True

headers_g = ["Câu", "Lệnh hỏi", "Đáp án (Đ/S)", "Câu", "Lệnh hỏi", "Đáp án (Đ/S)"]

def render_compact_grid(rows_data):
    t_grid = doc.add_table(rows=5, cols=6)
    t_grid.style = 'Table Grid'
    t_grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_grid.autofit = False
    for r in t_grid.rows:
        ExamDocxExporter._cant_split(r)

    for ci, h in enumerate(headers_g):
        c = t_grid.cell(0, ci)
        c.text = h
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0.5)
        p.paragraph_format.space_after = Pt(0.5)
        p.runs[0].font.name = "Times New Roman"
        p.runs[0].font.size = Pt(9.5)
        p.runs[0].bold = True
        ExamDocxExporter.set_cell_shading(c, "EAEAEA")

    for ri, (q1, q2) in enumerate(rows_data, 1):
        lbl = ["A", "B", "C", "D"][ri - 1]
        r_vals = [q1, lbl, "", q2, lbl, ""]
        for ci, val in enumerate(r_vals):
            c = t_grid.cell(ri, ci)
            c.text = val
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(0.5)
            if p.runs:
                p.runs[0].font.name = "Times New Roman"
                p.runs[0].font.size = Pt(9.5)
    ExamDocxExporter._set_col_widths(t_grid, [Inches(0.85), Inches(1.15), Inches(1.545), Inches(0.85), Inches(1.15), Inches(1.545)], total_width_dxa=10206)

render_compact_grid([("1", "2"), ("1", "2"), ("1", "2"), ("1", "2")])

p_gap = doc.add_paragraph()
p_gap.paragraph_format.space_before = Pt(1)
p_gap.paragraph_format.space_after = Pt(1)
p_gap.paragraph_format.keep_with_next = True

render_compact_grid([("3", "4"), ("3", "4"), ("3", "4"), ("3", "4")])

# 4 câu hỏi Đúng / Sai
for idx, tf_item in enumerate(exam_data["tfs"], 1):
    p_tf_q = doc.add_paragraph()
    p_tf_q.paragraph_format.space_before = Pt(1.5)
    p_tf_q.paragraph_format.space_after = Pt(0.5)
    p_tf_q.paragraph_format.line_spacing = 1.05
    p_tf_q.paragraph_format.keep_with_next = True

    r_tf_num = p_tf_q.add_run(f"Câu {idx}. “{tf_item['topic']}”")
    r_tf_num.font.name = "Times New Roman"
    r_tf_num.font.size = Pt(11)
    r_tf_num.bold = True

    for s_idx, stmt in enumerate(tf_item["statements"]):
        p_st = doc.add_paragraph()
        p_st.paragraph_format.space_before = Pt(0)
        p_st.paragraph_format.space_after = Pt(0.5)
        p_st.paragraph_format.left_indent = Inches(0.20)
        p_st.paragraph_format.line_spacing = 1.05
        if s_idx < len(tf_item["statements"]) - 1:
            p_st.paragraph_format.keep_with_next = True

        r_st_lbl = p_st.add_run(f"{stmt['label']}) ")
        r_st_lbl.font.name = "Times New Roman"
        r_st_lbl.font.size = Pt(11)
        r_st_lbl.bold = True

        r_st = p_st.add_run(stmt['content'])
        r_st.font.name = "Times New Roman"
        r_st.font.size = Pt(11)

# ==============================================================================
# TRANG 3: MẶT TRƯỚC TỜ 2 (TỰ LUẬN)
# ==============================================================================
doc.add_page_break()
ExamDocxExporter._render_header_box(doc, meta, code, is_essay=True)
ExamDocxExporter._render_cut_line(doc)
ExamDocxExporter._render_score_box(doc, code, is_essay=True)

for idx, es_q in enumerate(exam_data["essays"], 17):
    p_es = doc.add_paragraph()
    p_es.paragraph_format.space_before = Pt(2.5)
    p_es.paragraph_format.space_after = Pt(1.5)
    p_es.paragraph_format.line_spacing = 1.12
    p_es.paragraph_format.keep_with_next = True

    r_es_n = p_es.add_run(f"Câu {idx}. ({es_q['score']} điểm) ")
    r_es_n.font.name = "Times New Roman"
    r_es_n.font.size = Pt(12)
    r_es_n.bold = True

    r_es_c = p_es.add_run(es_q["content"])
    r_es_c.font.name = "Times New Roman"
    r_es_c.font.size = Pt(12)

p_end = doc.add_paragraph()
p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_end.paragraph_format.space_before = Pt(3)
p_end.paragraph_format.space_after = Pt(3)
p_end.paragraph_format.keep_with_next = True
r_end = p_end.add_run("------HẾT------")
r_end.bold = True
r_end.font.name = "Times New Roman"
r_end.font.size = Pt(11.5)

p_bl1 = doc.add_paragraph()
p_bl1.paragraph_format.space_before = Pt(2)
p_bl1.paragraph_format.space_after = Pt(1.5)
p_bl1.paragraph_format.keep_with_next = True
r_bl1 = p_bl1.add_run("BÀI LÀM")
r_bl1.bold = True
r_bl1.font.name = "Times New Roman"
r_bl1.font.size = Pt(12)

ExamDocxExporter._render_dot_lines(doc, count=16)

# ==============================================================================
# TRANG 4: MẶT SAU TỜ 2 (BÀI LÀM TIẾP THEO)
# ==============================================================================
doc.add_page_break()
ExamDocxExporter._render_x_cut_box(doc)

p_bl2 = doc.add_paragraph()
p_bl2.paragraph_format.space_before = Pt(2.5)
p_bl2.paragraph_format.space_after = Pt(1.5)
p_bl2.paragraph_format.keep_with_next = True
r_bl2 = p_bl2.add_run("BÀI LÀM (tiếp theo)")
r_bl2.bold = True
r_bl2.font.name = "Times New Roman"
r_bl2.font.size = Pt(12)

ExamDocxExporter._render_dot_lines(doc, count=27)

# Save
test_4p_path = os.path.abspath('scratch/test_single_exam_4pages.docx')
doc.save(test_4p_path)

# Verify with Word COM
word = win32com.client.Dispatch('Word.Application')
word.Visible = False
word.DisplayAlerts = 0
try:
    d = word.Documents.Open(test_4p_path, ReadOnly=True)
    pages = d.ComputeStatistics(2)
    print(f'>>> COMPUTED PAGES FOR SINGLE EXAM: {pages} <<<')
    for pg in range(1, pages + 1):
        rng = d.GoTo(1, 1, pg)
        rng2 = d.GoTo(1, 1, pg + 1) if pg < pages else d.Content
        end_pos = rng2.Start if pg < pages else d.Content.End
        rng.End = end_pos
        t = [l.strip() for l in rng.Text.split('\r') if l.strip()]
        first_line = t[0] if t else '<empty>'
        print(f'Page {pg} ({len(t):2d} lines): {first_line[:70]}')
    d.Close(False)
finally:
    word.Quit()
