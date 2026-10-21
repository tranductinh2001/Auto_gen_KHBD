import docx
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

path = 'BAO_CAO_DO_AN_CUOI_KY_CSDL_NANG_CAO_FINAL.docx'
doc = docx.Document(path)

print('File size:', os.path.getsize(path))
print('Sections count:', len(doc.sections))
for i, s in enumerate(doc.sections):
    print(f'Section {i}: start_type={s.start_type}, top={s.top_margin.inches:.2f}", left={s.left_margin.inches:.2f}"')
    h_text = ''.join(p.text for p in s.header.paragraphs).strip()
    f_text = ''.join(p.text for p in s.footer.paragraphs).strip()
    print(f'  Header: "{h_text[:40]}"')
    print(f'  Footer: "{f_text[:40]}"')

p_count = len(doc.paragraphs)
print('Total paragraphs:', p_count)
print('P0 (Cover):', doc.paragraphs[0].text[:60].replace('\n', ' '))
print('P1 (TOC Title):', doc.paragraphs[1].text)
print('P22 (Start of Body):', doc.paragraphs[22].text)

# Check first line indents of body paragraphs
indents = []
for i in range(23, min(35, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    if p.text.strip() and not p.text.strip().startswith(('1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '-', '•')):
        indents.append((p.text[:30], p.paragraph_format.first_line_indent.inches if p.paragraph_format.first_line_indent else 0))

print('Body indents sample:')
for t, ind in indents:
    print(f'  "{t}..." -> first_line_indent: {ind:.2f} in ({ind*2.54:.2f} cm)')
