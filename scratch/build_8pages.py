import os
import sys
import shutil
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import win32com.client

sys.stdout.reconfigure(encoding='utf-8')

BACKUP_PATH = 'BAO_CAO_DO_AN_CUOI_KY_CSDL_NANG_CAO_FINAL_BACKUP.docx'
TARGET_PATH = 'scratch/test_8pages.docx'

shutil.copy2(BACKUP_PATH, TARGET_PATH)
doc = docx.Document(TARGET_PATH)
print("Loaded clean document from backup.")

# 1. Section 0 Setup (Cover page)
sec0 = doc.sections[0]
sec0.page_width = Inches(8.27)
sec0.page_height = Inches(11.69)
sec0.top_margin = Inches(0.79)
sec0.bottom_margin = Inches(0.79)
sec0.left_margin = Inches(1.18)
sec0.right_margin = Inches(0.79)
sec0.different_first_page_header_footer = True

# Clear any header/footer on sec0
for h in [sec0.header, sec0.first_page_header]:
    for p in h.paragraphs:
        p.text = ""
for f in [sec0.footer, sec0.first_page_footer]:
    for p in f.paragraphs:
        p.text = ""

# In P0, remove any existing w:br so P0 doesn't add a double page break
p0 = doc.paragraphs[0]
for br in p0._p.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br'):
    if br.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type') == 'page':
        br.getparent().remove(br)

# We will create Section 1 for MỤC LỤC and Section 2 for BODY!
# How to create sections in python-docx:
# In Word XML:
# A section boundary is represented by `<w:sectPr>` inside the `<w:pPr>` of the last paragraph of that section!
# Let's create the TOC paragraphs before P1!

p1 = doc.paragraphs[1]

# Function to create paragraph before p1
def insert_p_before(ref_p, text="", align=WD_ALIGN_PARAGRAPH.LEFT, style_name="Normal"):
    new_p_elm = parse_xml(f'<w:p {nsdecls("w")}/>')
    ref_p._p.addprevious(new_p_elm)
    p_obj = docx.text.paragraph.Paragraph(new_p_elm, doc)
    p_obj.style = style_name
    p_obj.alignment = align
    if text:
        p_obj.text = text
    return p_obj

# Let's insert a page break / section break after P0:
# For Section 0 (Cover):
# We put a sectPr on p0's pPr
p0_pPr = p0._p.get_or_add_pPr()
sectPr0 = parse_xml(f'<w:sectPr {nsdecls("w")}>'
                    f'<w:pgSz w:w="11906" w:h="16838"/>' # 8.27 x 11.69 in
                    f'<w:pgMar w:top="1134" w:right="1134" w:bottom="1134" w:left="1701"/>'
                    f'<w:type w:val="nextPage"/>'
                    f'</w:sectPr>')
p0_pPr.append(sectPr0)

# NOW WE ARE IN SECTION 1: MỤC LỤC
# Title
p_toc_title = insert_p_before(p1, align=WD_ALIGN_PARAGRAPH.CENTER)
p_toc_title.paragraph_format.space_before = Pt(16)
p_toc_title.paragraph_format.space_after = Pt(14)
p_toc_title.paragraph_format.keep_with_next = True
r_tt = p_toc_title.add_run("MỤC LỤC")
r_tt.font.name = "Times New Roman"
r_tt.font.size = Pt(16)
r_tt.font.bold = True
r_tt.font.color.rgb = RGBColor(15, 44, 89)

# TOC items list: (Level, Text, Page)
toc_items = [
    (1, "1. THÔNG TIN TỔNG QUAN", "1"),
    (1, "2. BÀI TOÁN CẦN GIẢI QUYẾT VÀ Ý NGHĨA THỰC TIỄN", "1"),
    (2, "2.1. Thách thức trong quản lý dữ liệu giáo dục phân tán và phi cấu trúc", "1"),
    (2, "2.2. Ý nghĩa thực tiễn của đề tài", "1"),
    (1, "3. THIẾT KẾ CƠ SỞ DỮ LIỆU ĐỒ THỊ (PROPERTY GRAPH SCHEMA & DATA MODEL)", "1"),
    (2, "3.1. Tổng quan Kiến trúc Property Graph Model và Triết lý thiết kế", "1"),
    (2, "3.2. Sơ đồ Cấu trúc Đồ thị Tri thức Tổng thể (Meta-Graph Schema)", "2"),
    (2, "3.3. Đặc tả Danh mục Đỉnh (Node Labels) và Cạnh (Relationships)", "2"),
    (1, "4. NGUỒN DỮ LIỆU, QUY TRÌNH THU THẬP VÀ KẾT QUẢ THỰC NGHIỆM", "3"),
    (2, "4.1. Nguồn dữ liệu và Quy trình thu thập (ETL Pipeline)", "3"),
    (2, "4.2. Kết quả nạp dữ liệu thực nghiệm trong Neo4j Database 'edugraph'", "3"),
    (1, "5. KỸ THUẬT VÀ CÔNG CỤ GIẢI QUYẾT BÀI TOÁN", "3"),
    (1, "6. CÁC NỘI DUNG KHÁC: KIẾN TRÚC HỆ THỐNG, LUỒNG NGHIỆP VỤ VÀ MINH CHỨNG", "4"),
    (2, "6.1. Kiến trúc Web Dashboard và Phân hệ Nghiệp vụ Toàn diện", "4"),
    (2, "6.2. Các câu lệnh Cypher thực tế khai thác liên kết đồ thị", "4"),
    (2, "6.3. Hệ thống hình ảnh minh chứng thực tế trên hệ thống", "5"),
    (1, "7. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", "6"),
    (2, "7.1. Đánh giá kết quả đạt được", "6"),
    (2, "7.2. Hướng phát triển tiếp theo", "6"),
    (1, "8. TÀI LIỆU THAM KHẢO", "6"),
]

# Right tab stop at 16.0cm = 6.30 in = 9071 dxa
TOC_TAB_POS_DXA = 9071

last_toc_p = None
for lvl, text, pg in toc_items:
    p_item = insert_p_before(p1)
    p_item.paragraph_format.space_before = Pt(2.5 if lvl == 1 else 1)
    p_item.paragraph_format.space_after = Pt(2.5 if lvl == 1 else 1)
    p_item.paragraph_format.line_spacing = 1.15
    p_item.paragraph_format.left_indent = Inches(0 if lvl == 1 else 0.30)
    
    # Add tab stop with dot leader
    pPr = p_item._p.get_or_add_pPr()
    tabs_xml = f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="{TOC_TAB_POS_DXA}"/></w:tabs>'
    pPr.append(parse_xml(tabs_xml))

    r_t = p_item.add_run(text)
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(11.5 if lvl == 1 else 11)
    r_t.font.bold = (lvl == 1)
    if lvl == 1:
        r_t.font.color.rgb = RGBColor(15, 44, 89)
    else:
        r_t.font.color.rgb = RGBColor(30, 41, 59)

    r_tab = p_item.add_run("\t")
    r_tab.font.name = "Times New Roman"
    r_tab.font.size = Pt(11)

    r_pg = p_item.add_run(pg)
    r_pg.font.name = "Times New Roman"
    r_pg.font.size = Pt(11.5 if lvl == 1 else 11)
    r_pg.font.bold = (lvl == 1)
    r_pg.font.color.rgb = RGBColor(15, 44, 89) if lvl == 1 else RGBColor(30, 41, 59)
    
    last_toc_p = p_item

# At the end of Section 1 (Mục lục), add sectPr to boundary into Section 2 (Body)!
# Section 2 will have:
# - Page number restart at 1! (<w:pgNumType w:start="1"/>)
# - Running Header
# - Running Footer
p_last_toc_pPr = last_toc_p._p.get_or_add_pPr()
sectPr1 = parse_xml(f'<w:sectPr {nsdecls("w")}>'
                    f'<w:pgSz w:w="11906" w:h="16838"/>'
                    f'<w:pgMar w:top="1134" w:right="1134" w:bottom="1134" w:left="1701"/>'
                    f'<w:type w:val="nextPage"/>'
                    f'</w:sectPr>')
p_last_toc_pPr.append(sectPr1)

print("Inserted MỤC LỤC section successfully.")

# Save and inspect sections
doc.save(TARGET_PATH)
print("Saved temporary file. Re-opening to check sections...")
doc = docx.Document(TARGET_PATH)
print("Sections count now:", len(doc.sections))

# Section 2 is the Body!
sec_body = doc.sections[2]
# Configure pgNumType w:start="1" on sec_body
sec_body_pPr = sec_body._sectPr
pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1"/>')
sec_body_pPr.append(pgNumType)

# Unlink header and footer from Section 1 (Mục lục)
sec_body.header.is_linked_to_previous = False
sec_body.footer.is_linked_to_previous = False

# Configure running header on sec_body
hp = sec_body.header.paragraphs[0]
hp.text = ""
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.paragraph_format.space_before = Pt(0)
hp.paragraph_format.space_after = Pt(4)
hrun = hp.add_run("TRƯỜNG ĐẠI HỌC QUY NHƠN – BÁO CÁO ĐỀ ÁN MÔN CƠ SỞ DỮ LIỆU NÂNG CAO")
hrun.font.name = "Times New Roman"
hrun.font.size = Pt(9)
hrun.font.italic = True
hrun.font.color.rgb = RGBColor(100, 116, 139)

pPr = hp._p.get_or_add_pPr()
pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="2" w:color="CBD5E1"/></w:pBdr>')
pPr.append(pBdr)

# Configure running footer on sec_body
fp = sec_body.footer.paragraphs[0]
fp.text = ""
fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
fp.paragraph_format.space_before = Pt(4)
fp.paragraph_format.space_after = Pt(0)

frun_left = fp.add_run("Học viên: Trần Đức Tình – MSHV: 8291581013")
frun_left.font.name = "Times New Roman"
frun_left.font.size = Pt(9.5)
frun_left.font.color.rgb = RGBColor(100, 116, 139)

fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.30), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)
fp.add_run("\t")

frun_pg = fp.add_run("Trang ")
frun_pg.font.name = "Times New Roman"
frun_pg.font.size = Pt(9.5)
frun_pg.font.color.rgb = RGBColor(100, 116, 139)

fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
fp._p.append(fldSimple)

# Clear headers/footers on Section 0 and Section 1
doc.sections[0].header.is_linked_to_previous = False
doc.sections[0].footer.is_linked_to_previous = False
doc.sections[0].header.paragraphs[0].text = ""
doc.sections[0].footer.paragraphs[0].text = ""

doc.sections[1].header.is_linked_to_previous = False
doc.sections[1].footer.is_linked_to_previous = False
doc.sections[1].header.paragraphs[0].text = ""
doc.sections[1].footer.paragraphs[0].text = ""

doc.save(TARGET_PATH)
print("Configured 3 sections, headers, and footers successfully.")
