import os
import sys
import shutil
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

BACKUP_PATH = 'BAO_CAO_DO_AN_CUOI_KY_CSDL_NANG_CAO_FINAL_BACKUP.docx'
FINAL_PATH = 'BAO_CAO_DO_AN_CUOI_KY_CSDL_NANG_CAO_FINAL.docx'

if not os.path.exists(BACKUP_PATH):
    raise FileNotFoundError(f"Backup file {BACKUP_PATH} does not exist!")

print(f"Reading from backup: {BACKUP_PATH}")
shutil.copy2(BACKUP_PATH, FINAL_PATH)
doc = docx.Document(FINAL_PATH)
print("Loaded document successfully.")

# ==============================================================================
# 1. PAGE SETUP & MARGINS (A4: Top 1.8cm, Bottom 1.8cm, Left 3.0cm, Right 1.9cm)
# ==============================================================================
for sec in doc.sections:
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.71)    # 1.8 cm
    sec.bottom_margin = Inches(0.71) # 1.8 cm
    sec.left_margin = Inches(1.18)   # 3.0 cm (gáy chuẩn luận văn)
    sec.right_margin = Inches(0.75)  # 1.9 cm

# In P0 (Cover page paragraph), remove any existing w:br page breaks
p0 = doc.paragraphs[0]
for br in p0._p.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br'):
    if br.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type') == 'page':
        br.getparent().remove(br)

# Create Section 0 boundary on P0 (Bìa -> Section break sang Mục lục)
p0_pPr = p0._p.get_or_add_pPr()
sectPr0 = parse_xml(f'<w:sectPr {nsdecls("w")}>'
                    f'<w:pgSz w:w="11906" w:h="16838"/>'
                    f'<w:pgMar w:top="1020" w:right="1077" w:bottom="1020" w:left="1701"/>'
                    f'<w:type w:val="nextPage"/>'
                    f'</w:sectPr>')
p0_pPr.append(sectPr0)

# ==============================================================================
# 2. SECTION 1: MỤC LỤC
# ==============================================================================
p1 = doc.paragraphs[1]

def insert_p_before(ref_p, text="", align=WD_ALIGN_PARAGRAPH.LEFT, style_name="Normal"):
    new_p_elm = parse_xml(f'<w:p {nsdecls("w")}/>')
    ref_p._p.addprevious(new_p_elm)
    p_obj = docx.text.paragraph.Paragraph(new_p_elm, doc)
    p_obj.style = style_name
    p_obj.alignment = align
    if text:
        p_obj.text = text
    return p_obj

# Tiêu đề Mục Lục
p_toc_title = insert_p_before(p1, align=WD_ALIGN_PARAGRAPH.CENTER)
p_toc_title.paragraph_format.space_before = Pt(12)
p_toc_title.paragraph_format.space_after = Pt(10)
p_toc_title.paragraph_format.keep_with_next = True
r_tt = p_toc_title.add_run("MỤC LỤC")
r_tt.font.name = "Times New Roman"
r_tt.font.size = Pt(15)
r_tt.font.bold = True
r_tt.font.color.rgb = RGBColor(15, 44, 89)

# Các mục mục lục với số trang hiển thị (bắt đầu từ Trang 1 sau mục lục)
toc_items = [
    (1, "1. THÔNG TIN TỔNG QUAN", "1"),
    (1, "2. BÀI TOÁN CẦN GIẢI QUYẾT VÀ Ý NGHĨA THỰC TIỄN", "1"),
    (2, "2.1. Thách thức trong quản lý dữ liệu giáo dục phân tán và phi cấu trúc", "1"),
    (2, "2.2. Ý nghĩa thực tiễn của đề tài", "1"),
    (1, "3. THIẾT KẾ CƠ SỞ DỮ LIỆU ĐỒ THỊ (PROPERTY GRAPH SCHEMA & DATA MODEL)", "1"),
    (2, "3.1. Tổng quan Kiến trúc Property Graph Model và Triết lý thiết kế", "1"),
    (2, "3.2. Sơ đồ Cấu trúc Đồ thị Tri thức Tổng thể (Meta-Graph Schema)", "2"),
    (2, "3.3. Đặc tả Danh mục Đỉnh (Node Labels) và Cạnh (Relationships)", "2"),
    (1, "4. NGUỒN DỮ LIỆU, QUY TRÌNH THU THẬP VÀ KẾT QUẢ THỰC NGHIỆM", "3"),
    (2, "4.1. Nguồn dữ liệu và Quy trình thu thập (ETL Pipeline)", "3"),
    (2, "4.2. Kết quả nạp dữ liệu thực nghiệm trong Neo4j Database 'edugraph'", "3"),
    (1, "5. KỸ THUẬT VÀ CÔNG CỤ GIẢI QUYẾT BÀI TOÁN", "3"),
    (1, "6. CÁC NỘI DUNG KHÁC: KIẾN TRÚC HỆ THỐNG, LUỒNG NGHIỆP VỤ VÀ MINH CHỨNG", "4"),
    (2, "6.1. Kiến trúc Web Dashboard và Phân hệ Nghiệp vụ Toàn diện", "4"),
    (2, "6.2. Các câu lệnh Cypher thực tế khai thác liên kết đồ thị", "4"),
    (2, "6.3. Hệ thống hình ảnh minh chứng thực tế trên hệ thống", "4"),
    (1, "7. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", "6"),
    (2, "7.1. Đánh giá kết quả đạt được", "6"),
    (2, "7.2. Hướng phát triển tiếp theo", "6"),
    (1, "8. TÀI LIỆU THAM KHẢO", "6"),
]

TOC_TAB_POS_DXA = 9071 # Căn sát lề phải
last_toc_p = None
for lvl, text, pg in toc_items:
    p_item = insert_p_before(p1)
    p_item.paragraph_format.space_before = Pt(1.5 if lvl == 1 else 0.5)
    p_item.paragraph_format.space_after = Pt(1.5 if lvl == 1 else 0.5)
    p_item.paragraph_format.line_spacing = 1.10
    p_item.paragraph_format.left_indent = Inches(0 if lvl == 1 else 0.28)
    
    pPr = p_item._p.get_or_add_pPr()
    tabs_xml = f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="{TOC_TAB_POS_DXA}"/></w:tabs>'
    pPr.append(parse_xml(tabs_xml))

    r_t = p_item.add_run(text)
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(11.5 if lvl == 1 else 10.5)
    r_t.font.bold = (lvl == 1)
    r_t.font.color.rgb = RGBColor(15, 44, 89) if lvl == 1 else RGBColor(30, 41, 59)

    r_tab = p_item.add_run("\t")
    r_tab.font.name = "Times New Roman"
    r_tab.font.size = Pt(10.5)

    r_pg = p_item.add_run(pg)
    r_pg.font.name = "Times New Roman"
    r_pg.font.size = Pt(11.5 if lvl == 1 else 10.5)
    r_pg.font.bold = (lvl == 1)
    r_pg.font.color.rgb = RGBColor(15, 44, 89) if lvl == 1 else RGBColor(30, 41, 59)
    
    last_toc_p = p_item

# Boundary Section 1 (Mục lục) -> Section 2 (Thân bài bắt đầu từ "1. THÔNG TIN TỔNG QUAN")
p_last_toc_pPr = last_toc_p._p.get_or_add_pPr()
sectPr1 = parse_xml(f'<w:sectPr {nsdecls("w")}>'
                    f'<w:pgSz w:w="11906" w:h="16838"/>'
                    f'<w:pgMar w:top="1020" w:right="1077" w:bottom="1020" w:left="1701"/>'
                    f'<w:type w:val="nextPage"/>'
                    f'</w:sectPr>')
p_last_toc_pPr.append(sectPr1)

# Lưu và mở lại để python-docx cập nhật chuẩn xác cấu trúc 3 sections
doc.save(FINAL_PATH)
doc = docx.Document(FINAL_PATH)

# ==============================================================================
# 3. HEADER & FOOTER CONFIGURATION (ĐÁNH SỐ TRANG SAU MỤC LỤC)
# ==============================================================================
# Section 0 (Bìa): Không có Header, không có Footer
doc.sections[0].header.is_linked_to_previous = False
doc.sections[0].footer.is_linked_to_previous = False
doc.sections[0].different_first_page_header_footer = False
doc.sections[0].header.paragraphs[0].text = ""
doc.sections[0].footer.paragraphs[0].text = ""

# Section 1 (Mục lục): Không có Header, không có Footer
doc.sections[1].header.is_linked_to_previous = False
doc.sections[1].footer.is_linked_to_previous = False
doc.sections[1].different_first_page_header_footer = False
doc.sections[1].header.paragraphs[0].text = ""
doc.sections[1].footer.paragraphs[0].text = ""

# Section 2 (Toàn bộ Nội dung thân bài)
sec_body = doc.sections[2]
sec_body.different_first_page_header_footer = False
sec_body.header.is_linked_to_previous = False
sec_body.footer.is_linked_to_previous = False

# Thiết lập khởi đầu số trang từ 1 ngay sau Mục lục
sec_body_pPr = sec_body._sectPr
# Xóa pgNumType cũ nếu có
for c in list(sec_body_pPr):
    if c.tag.endswith('pgNumType'):
        sec_body_pPr.remove(c)
pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1"/>')
sec_body_pPr.append(pgNumType)

# Header thân bài (Góc phải có gạch phân cách thanh lịch)
hp = sec_body.header.paragraphs[0]
hp.text = ""
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.paragraph_format.space_before = Pt(0)
hp.paragraph_format.space_after = Pt(2)

hrun = hp.add_run("TRƯỜNG ĐẠI HỌC QUY NHƠN – BÁO CÁO ĐỀ ÁN MÔN CƠ SỞ DỮ LIỆU NÂNG CAO")
hrun.font.name = "Times New Roman"
hrun.font.size = Pt(8.5)
hrun.font.italic = True
hrun.font.color.rgb = RGBColor(100, 116, 139)

pPr_h = hp._p.get_or_add_pPr()
pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="2" w:color="CBD5E1"/></w:pBdr>')
pPr_h.append(pBdr)

# Footer thân bài: Trái là Tên học viên, Phải là "Trang X"
fp = sec_body.footer.paragraphs[0]
fp.text = ""
fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
fp.paragraph_format.space_before = Pt(2)
fp.paragraph_format.space_after = Pt(0)

frun_left = fp.add_run("Học viên: Trần Đức Tình – MSHV: 8291581013")
frun_left.font.name = "Times New Roman"
frun_left.font.size = Pt(9.0)
frun_left.font.color.rgb = RGBColor(100, 116, 139)

fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.30), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)
fp.add_run("\t")

frun_pg = fp.add_run("Trang ")
frun_pg.font.name = "Times New Roman"
frun_pg.font.size = Pt(9.0)
frun_pg.font.color.rgb = RGBColor(100, 116, 139)

fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
fp._p.append(fldSimple)

# ==============================================================================
# 4. ĐỊNH DẠNG BẢNG BIỂU CHUYÊN NGHIỆP
# ==============================================================================
def set_cell_shd(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    for c in list(tcPr):
        if c.tag.endswith('shd'): tcPr.remove(c)
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))

def set_cell_pad(cell, t=30, b=30, l=80, r=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{t}" w:type="dxa"/><w:bottom w:w="{b}" w:type="dxa"/><w:left w:w="{l}" w:type="dxa"/><w:right w:w="{r}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_tbl_borders(table):
    tblPr = table._tbl.tblPr
    for c in list(tblPr):
        if c.tag.endswith('tblBorders'): tblPr.remove(c)
    tblPr.append(parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="6" w:space="0" w:color="1E3A8A"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="1E3A8A"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:insideV w:val="none"/></w:tblBorders>'))

def format_table(table, col_widths, center_cols=[0]):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_tbl_borders(table)

    # Header Row
    table.rows[0]._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    table.rows[0]._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    for cell in table.rows[0].cells:
        set_cell_shd(cell, "1E3A8A")
        set_cell_pad(cell, t=50, b=50, l=80, r=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0.5)
        p.paragraph_format.space_after = Pt(0.5)
        p.paragraph_format.line_spacing = 1.05
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(8.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for ri, row in enumerate(table.rows[1:], 1):
        row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg = "FFFFFF" if ri % 2 == 1 else "F8FAFC"
        is_total = "TỔNG CỘNG" in row.cells[0].text

        for ci, cell in enumerate(row.cells):
            set_cell_shd(cell, bg if not is_total else "EEF2F6")
            set_cell_pad(cell, t=30, b=30, l=80, r=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(0.5)
            p.paragraph_format.line_spacing = 1.05
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(8.0)
                if is_total:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(15, 23, 42)
                else:
                    r.font.color.rgb = RGBColor(30, 41, 59)

    total_w = sum(col_widths)
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{total_w}" w:type="dxa"/>'))
    else:
        tblW.attrib[qn('w:w')] = str(total_w)
        tblW.attrib[qn('w:type')] = 'dxa'

    for row in table.rows:
        for ci, w in enumerate(col_widths):
            if ci < len(row.cells):
                tcPr = row.cells[ci]._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn('w:tcW'))
                if tcW is None:
                    tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{w}" w:type="dxa"/>'))
                else:
                    tcW.attrib[qn('w:w')] = str(w)

format_table(doc.tables[0], [600, 1500, 2500, 2671, 1800], center_cols=[0])
format_table(doc.tables[1], [600, 2000, 2300, 900, 3271], center_cols=[0, 3])
format_table(doc.tables[2], [1271, 1300, 1300, 1300, 1300, 1300, 1300], center_cols=[0, 1, 2, 3, 4, 5, 6])
format_table(doc.tables[3], [2271, 2300, 4500], center_cols=[])

# Table 4 (Hộp mã Cypher)
t_code = doc.tables[4]
t_code.alignment = WD_TABLE_ALIGNMENT.CENTER
t_code.autofit = False
c_code = t_code.rows[0].cells[0]
set_cell_shd(c_code, "F8FAFC")
set_cell_pad(c_code, t=60, b=60, l=120, r=100)
tcPr = c_code._tc.get_or_add_tcPr()
tcPr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/></w:tcBorders>'))
c_code.width = Inches(6.30)
tcW = tcPr.find(qn('w:tcW'))
if tcW is not None: tcW.attrib[qn('w:w')] = "9071"

for p in c_code.paragraphs:
    p.paragraph_format.space_before = Pt(0.5)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.line_spacing = 1.05
    for r in p.runs:
        r.font.name = "Consolas"
        r.font.size = Pt(7.0)
        if r.text.strip().startswith("//"):
            r.font.color.rgb = RGBColor(100, 116, 139)
            r.font.italic = True
        else:
            r.font.color.rgb = RGBColor(15, 23, 42)

# ==============================================================================
# 5. SCALE HÌNH ẢNH MINH CHỨNG (ĐỂ VỪA VẶN TRANG VÀ GIỮ NÉT)
# ==============================================================================
def scale_draw(draw_elm, max_w_in=None, max_h_in=None):
    ext = draw_elm.find('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent')
    if ext is None: return
    cx = int(ext.attrib.get('cx', 0))
    cy = int(ext.attrib.get('cy', 0))
    if cx <= 0 or cy <= 0: return

    cur_w_in = cx / 914400.0
    cur_h_in = cy / 914400.0

    target_w = cur_w_in
    target_h = cur_h_in

    if max_w_in and target_w > max_w_in:
        ratio = max_w_in / target_w
        target_w = max_w_in
        target_h = target_h * ratio

    if max_h_in and target_h > max_h_in:
        ratio = max_h_in / target_h
        target_h = max_h_in
        target_w = target_w * ratio

    ext.attrib['cx'] = str(int(target_w * 914400))
    ext.attrib['cy'] = str(int(target_h * 914400))

for i, p in enumerate(doc.paragraphs):
    if i == 0: continue
    draws = p._p.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
    if not draws: continue

    if '3.2. Sơ đồ' in p.text:
        for d in draws: scale_draw(d, max_w_in=4.60, max_h_in=2.00)
    elif i <= 72: # Hình 1
        for d in draws: scale_draw(d, max_w_in=4.20, max_h_in=1.60)
    elif i <= 78: # Hình 2, 3, 4
        for d in draws: scale_draw(d, max_w_in=4.00, max_h_in=1.40)
    elif i <= 82: # Hình 5, 6
        for d in draws: scale_draw(d, max_w_in=4.00, max_h_in=1.40)
    elif i >= 83: # Hình 7 (P85 - 2 drawings)
        for d in draws: scale_draw(d, max_w_in=3.80, max_h_in=1.05)

# ==============================================================================
# 6. ĐỊNH DẠNG ĐOẠN VĂN: THỤT ĐẦU DÒNG, CĂN ĐỀU, GIÃN CÁCH CHUẨN ĐẸP
# ==============================================================================
H1_REGEX = re.compile(r'^[1-8]\.\s+([A-ZÀ-Ỹ0-9\s:()\'/&,-]+)$')
H2_REGEX = re.compile(r'^[1-8]\.[1-9]\.\s+')

for i, p in enumerate(doc.paragraphs):
    # Bỏ qua bìa và các đoạn thuộc Mục Lục
    if i == 0 or i <= len(toc_items) + 1:
        continue

    t = p.text.strip()
    draws = p._p.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')

    # Paragraph chứa ảnh
    if draws and (not t or len(t) < 3):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        continue

    # Dòng trống
    if not t:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs: r.font.size = Pt(1)
        continue

    # Tiêu đề Cấp 1 (Heading 1)
    if H1_REGEX.match(t):
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.10
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Pt(0)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 44, 89)

    # Tiêu đề Cấp 2 (Heading 2)
    elif H2_REGEX.match(t):
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.10
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Pt(0)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(30, 58, 138)

    # Chú thích hình (Figure Caption)
    elif t.startswith('Hình '):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3.5)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        colon_idx = t.find(':')
        if colon_idx != -1:
            p.text = ""
            r_lbl = p.add_run(t[:colon_idx+1])
            r_lbl.font.name = "Times New Roman"
            r_lbl.font.size = Pt(9.0)
            r_lbl.font.bold = True
            r_lbl.font.color.rgb = RGBColor(51, 65, 85)
            r_desc = p.add_run(t[colon_idx+1:])
            r_desc.font.name = "Times New Roman"
            r_desc.font.size = Pt(9.0)
            r_desc.font.italic = True
            r_desc.font.color.rgb = RGBColor(51, 65, 85)

    # Tài liệu tham khảo (Mục 8)
    elif i >= 90 and re.match(r'^[1-6]\.\s+', t):
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0.5)
        p.paragraph_format.space_after = Pt(1.0)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.0)
            r.font.color.rgb = RGBColor(15, 23, 42)

    # Các ý gạch đầu dòng (Bullet points)
    elif p.style.name == 'List Bullet' or t.startswith(('- ', '• ', '1. ', '2. ', '3. ', '4. ')):
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0.5)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.12
        p.paragraph_format.left_indent = Inches(0.30)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        if ':' in t:
            parts = t.split(':', 1)
            p.text = ""
            r_key = p.add_run(parts[0] + ":")
            r_key.font.name = "Times New Roman"
            r_key.font.size = Pt(11.0)
            r_key.font.bold = True
            r_key.font.color.rgb = RGBColor(15, 23, 42)

            r_val = p.add_run(parts[1])
            r_val.font.name = "Times New Roman"
            r_val.font.size = Pt(11.0)
            r_val.font.color.rgb = RGBColor(30, 41, 59)
        else:
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11.0)

    # Đoạn văn bản thông thường (Body text): Thụt đầu dòng 0.35 inch (~0.9 cm) chuẩn tiếng Việt
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.first_line_indent = Inches(0.35) # Thụt đầu dòng đẹp mắt
        p.paragraph_format.left_indent = Pt(0)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(11.0)
            r.font.color.rgb = RGBColor(15, 23, 42)

doc.save(FINAL_PATH)
print(f"SUCCESS! Outputted formatted document directly to: {FINAL_PATH}")
