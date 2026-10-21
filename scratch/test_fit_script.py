import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client
import os, sys

sys.path.insert(0, os.getcwd())
sys.stdout.reconfigure(encoding='utf-8')

from core.exam_builder import ExamMatrixEngine, ExamDocxExporter

engine = ExamMatrixEngine(
    grade_level='Lớp 9',
    exam_title='KIỂM TRA GIỮA HỌC KỲ I',
    school_name='TRƯỜNG TH&THCS ĐĂK TRÔI',
    governing_body='UBND XÃ KON CHIÊNG',
    school_year='2026 - 2027',
    teacher_name='Lê Thị Phương Thảo',
    head_name='Trần Thị Tuyết Ngân'
)
dossier = engine.generate_exam_dossier()
exam_data = dossier['th01']
meta = dossier['meta']
code = exam_data['code']

# Test single complete 4-page exam for TH01
doc = docx.Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.PORTRAIT
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(0.59)
sec.bottom_margin = Inches(0.59)
sec.left_margin = Inches(0.59)
sec.right_margin = Inches(0.59)

# ==============================================================================
# TRANG 1: MẶT TRƯỚC TỜ 1 (TRẮC NGHIỆM CÂU 1 ĐẾN CÂU 7)
# ==============================================================================
ExamDocxExporter._render_header_box(doc, meta, code, is_essay=False)
ExamDocxExporter._render_cut_line(doc)
ExamDocxExporter._render_score_box(doc, code, is_essay=False)

p_p1 = doc.add_paragraph()
p_p1.paragraph_format.space_before = Pt(2.0)
p_p1.paragraph_format.space_after = Pt(1.0)
p_p1.paragraph_format.keep_with_next = True
r = p_p1.add_run('I. Khoanh tròn vào chữ cái đứng trước câu trả lời đúng (3,0 điểm)')
r.font.name = 'Times New Roman'
r.font.size = Pt(11.5)
r.bold = True

def render_mcq(q, idx, fs=11):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(1.5)
    p_q.paragraph_format.space_after = Pt(0.5)
    p_q.paragraph_format.line_spacing = 1.05
    p_q.paragraph_format.keep_with_next = True
    rq = p_q.add_run(f'Câu {idx}. ')
    rq.font.name = 'Times New Roman'
    rq.font.size = Pt(fs)
    rq.bold = True
    rc = p_q.add_run(q['content'])
    rc.font.name = 'Times New Roman'
    rc.font.size = Pt(fs)
    
    opts = q['options']
    opt_keys = ['A', 'B', 'C', 'D']
    max_opt_len = max(len(str(opts.get(k, ''))) for k in opt_keys)
    
    if max_opt_len <= 14:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.20)
        p.paragraph_format.space_before = Pt(0.5)
        p.paragraph_format.space_after = Pt(0.8)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.tab_stops.add_tab_stop(Inches(1.85), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.60), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(5.35), WD_TAB_ALIGNMENT.LEFT)
        for i, k in enumerate(opt_keys):
            if i > 0: p.add_run('\t')
            p.add_run(f'{k}. ').bold = True
            p.add_run(str(opts.get(k, '')))
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(fs)
    elif max_opt_len <= 56:
        for p_idx, (k1, k2) in enumerate([('A', 'B'), ('C', 'D')]):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.20)
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(0.8)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.keep_with_next = (p_idx == 0)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(3.60), WD_TAB_ALIGNMENT.LEFT)
            p.add_run(f'{k1}. ').bold = True
            p.add_run(str(opts.get(k1, '')))
            p.add_run('\t')
            p.add_run(f'{k2}. ').bold = True
            p.add_run(str(opts.get(k2, '')))
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(fs)
    else:
        for i, k in enumerate(opt_keys):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.20)
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(0.8)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.keep_with_next = (i < 3)
            p.add_run(f'{k}. ').bold = True
            p.add_run(str(opts.get(k, '')))
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(fs)

P1_COUNT = 7
for idx, q in enumerate(exam_data['mcqs'][:P1_COUNT], 1):
    render_mcq(q, idx, fs=11)

# ==============================================================================
# TRANG 2: MẶT SAU TỜ 1 (TRẮC NGHIỆM CÂU 8..12 + PHẦN II ĐÚNG/SAI)
# ==============================================================================
doc.add_page_break()

# Render X-box compact (1500 dxa ~ 75 pt)
def render_compact_x_box(d):
    t_box = d.add_table(rows=1, cols=1)
    t_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_box.autofit = False
    ExamDocxExporter._cant_split(t_box.rows[0])
    ExamDocxExporter._set_row_height_exact(t_box.rows[0], 1500)
    c = t_box.cell(0, 0)
    c.width = Inches(7.09)
    ExamDocxExporter.set_cell_border(
        c,
        top={'val': 'single', 'sz': 4, 'space': 0, 'color': 'auto'},
        left={'val': 'single', 'sz': 4, 'space': 0, 'color': 'auto'},
        bottom={'val': 'single', 'sz': 4, 'space': 0, 'color': 'auto'},
        right={'val': 'single', 'sz': 4, 'space': 0, 'color': 'auto'},
        tl2br={'val': 'single', 'sz': 4, 'space': 0, 'color': 'auto'},
        tr2bl={'val': 'single', 'sz': 4, 'space': 0, 'color': 'auto'}
    )
    p_in = c.paragraphs[0]
    p_in.paragraph_format.space_before = Pt(0)
    p_in.paragraph_format.space_after = Pt(0)
    r_in = p_in.add_run()
    r_in.font.size = Pt(4)
    ExamDocxExporter._set_col_widths(t_box, [Inches(7.09)], total_width_dxa=10206)
    ExamDocxExporter._render_cut_line(d)

render_compact_x_box(doc)

def render_mcq_p2(q, idx):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(1.0)
    p_q.paragraph_format.space_after = Pt(0.2)
    p_q.paragraph_format.line_spacing = 1.02
    p_q.paragraph_format.keep_with_next = True
    rq = p_q.add_run(f'Câu {idx}. ')
    rq.font.name = 'Times New Roman'
    rq.font.size = Pt(10)
    rq.bold = True
    rc = p_q.add_run(q['content'])
    rc.font.name = 'Times New Roman'
    rc.font.size = Pt(10)
    
    opts = q['options']
    opt_keys = ['A', 'B', 'C', 'D']
    max_opt_len = max(len(str(opts.get(k, ''))) for k in opt_keys)
    
    if max_opt_len <= 14:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.20)
        p.paragraph_format.space_before = Pt(0.2)
        p.paragraph_format.space_after = Pt(0.4)
        p.paragraph_format.line_spacing = 1.02
        p.paragraph_format.tab_stops.add_tab_stop(Inches(1.85), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.60), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(5.35), WD_TAB_ALIGNMENT.LEFT)
        for i, k in enumerate(opt_keys):
            if i > 0: p.add_run('\t')
            p.add_run(f'{k}. ').bold = True
            p.add_run(str(opts.get(k, '')))
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
    elif max_opt_len <= 56:
        for p_idx, (k1, k2) in enumerate([('A', 'B'), ('C', 'D')]):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.20)
            p.paragraph_format.space_before = Pt(0.2)
            p.paragraph_format.space_after = Pt(0.4)
            p.paragraph_format.line_spacing = 1.02
            p.paragraph_format.keep_with_next = (p_idx == 0)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(3.60), WD_TAB_ALIGNMENT.LEFT)
            p.add_run(f'{k1}. ').bold = True
            p.add_run(str(opts.get(k1, '')))
            p.add_run('\t')
            p.add_run(f'{k2}. ').bold = True
            p.add_run(str(opts.get(k2, '')))
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
    else:
        for i, k in enumerate(opt_keys):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.20)
            p.paragraph_format.space_before = Pt(0.2)
            p.paragraph_format.space_after = Pt(0.4)
            p.paragraph_format.line_spacing = 1.02
            p.paragraph_format.keep_with_next = (i < 3)
            p.add_run(f'{k}. ').bold = True
            p.add_run(str(opts.get(k, '')))
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)

for idx, q in enumerate(exam_data['mcqs'][P1_COUNT:], P1_COUNT + 1):
    render_mcq_p2(q, idx)

p_p2 = doc.add_paragraph()
p_p2.paragraph_format.space_before = Pt(1.0)
p_p2.paragraph_format.space_after = Pt(0.5)
p_p2.paragraph_format.keep_with_next = True
r2 = p_p2.add_run('II. Viết chữ Đ(đúng) hoặc S(sai) vào các câu trả lời tương ứng trong các câu sau: (4,0 điểm)')
r2.font.name = 'Times New Roman'
r2.font.size = Pt(10.5)
r2.bold = True

headers_12 = ['Câu', 'Lệnh', 'Đ/S', 'Câu', 'Lệnh', 'Đ/S', 'Câu', 'Lệnh', 'Đ/S', 'Câu', 'Lệnh', 'Đ/S']
t_grid12 = doc.add_table(rows=5, cols=12)
t_grid12.style = 'Table Grid'
t_grid12.alignment = WD_TABLE_ALIGNMENT.CENTER
tblPr = t_grid12._tbl.tblPr
tblCellMar = parse_xml(f'<w:tblCellMar {nsdecls("w")}><w:top w:w="10" w:type="dxa"/><w:bottom w:w="10" w:type="dxa"/><w:left w:w="25" w:type="dxa"/><w:right w:w="25" w:type="dxa"/></w:tblCellMar>')
tblPr.append(tblCellMar)

for r in t_grid12.rows:
    ExamDocxExporter._cant_split(r)

for ci, h in enumerate(headers_12):
    c = t_grid12.cell(0, ci)
    c.text = h
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.runs[0].font.name = 'Times New Roman'
    p.runs[0].font.size = Pt(8.5)
    p.runs[0].bold = True
    ExamDocxExporter.set_cell_shading(c, 'EAEAEA')

for ri, lbl in enumerate(['A', 'B', 'C', 'D'], 1):
    r_vals = ['1', lbl, '', '2', lbl, '', '3', lbl, '', '4', lbl, '']
    for ci, val in enumerate(r_vals):
        c = t_grid12.cell(ri, ci)
        c.text = val
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if p.runs:
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(8.5)

col_w = [Inches(0.42), Inches(0.42), Inches(0.85)] * 4
ExamDocxExporter._set_col_widths(t_grid12, col_w, total_width_dxa=10206)

# 4 True/False questions
for idx, tf_item in enumerate(exam_data['tfs'], 1):
    p_tf_q = doc.add_paragraph()
    p_tf_q.paragraph_format.space_before = Pt(0.8)
    p_tf_q.paragraph_format.space_after = Pt(0)
    p_tf_q.paragraph_format.line_spacing = 1.0
    p_tf_q.paragraph_format.keep_with_next = True
    top_txt = tf_item['topic'][:65]
    r_num = p_tf_q.add_run(f'Câu {idx}. “{top_txt}”')
    r_num.font.name = 'Times New Roman'
    r_num.font.size = Pt(9.5)
    r_num.bold = True
    for s_idx, stmt in enumerate(tf_item['statements']):
        p_st = doc.add_paragraph()
        p_st.paragraph_format.space_before = Pt(0)
        p_st.paragraph_format.space_after = Pt(0.1)
        p_st.paragraph_format.left_indent = Inches(0.20)
        p_st.paragraph_format.line_spacing = 1.0
        if s_idx < 3: p_st.paragraph_format.keep_with_next = True
        lbl_val = stmt['label']
        r_lbl = p_st.add_run(f'{lbl_val}) ')
        r_lbl.bold = True
        r_st = p_st.add_run(stmt['content'])
        for r in [r_lbl, r_st]:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.5)

# ==============================================================================
# TRANG 3: MẶT TRƯỚC TỜ 2 (TỰ LUẬN)
# ==============================================================================
doc.add_page_break()
ExamDocxExporter._render_header_box(doc, meta, code, is_essay=True)
ExamDocxExporter._render_cut_line(doc)
ExamDocxExporter._render_score_box(doc, code, is_essay=True)

for idx, es_q in enumerate(exam_data['essays'], 17):
    p_es = doc.add_paragraph()
    p_es.paragraph_format.space_before = Pt(2.0)
    p_es.paragraph_format.space_after = Pt(1.0)
    p_es.paragraph_format.line_spacing = 1.08
    p_es.paragraph_format.keep_with_next = True

    r_es_n = p_es.add_run(f"Câu {idx}. ({es_q['score']} điểm) ")
    r_es_n.font.name = "Times New Roman"
    r_es_n.font.size = Pt(11.5)
    r_es_n.bold = True

    r_es_c = p_es.add_run(es_q['content'])
    r_es_c.font.name = "Times New Roman"
    r_es_c.font.size = Pt(11.5)

p_end = doc.add_paragraph()
p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_end.paragraph_format.space_before = Pt(2)
p_end.paragraph_format.space_after = Pt(2)
p_end.paragraph_format.keep_with_next = True
r_end = p_end.add_run("------HẾT------")
r_end.bold = True
r_end.font.name = "Times New Roman"
r_end.font.size = Pt(11)

p_bl1 = doc.add_paragraph()
p_bl1.paragraph_format.space_before = Pt(2)
p_bl1.paragraph_format.space_after = Pt(1)
p_bl1.paragraph_format.keep_with_next = True
r_bl1 = p_bl1.add_run("BÀI LÀM")
r_bl1.bold = True
r_bl1.font.name = "Times New Roman"
r_bl1.font.size = Pt(11.5)

ExamDocxExporter._render_dot_lines(doc, count=16)

# ==============================================================================
# TRANG 4: MẶT SAU TỜ 2 (BÀI LÀM TIẾP THEO)
# ==============================================================================
doc.add_page_break()
render_compact_x_box(doc)

p_bl2 = doc.add_paragraph()
p_bl2.paragraph_format.space_before = Pt(2)
p_bl2.paragraph_format.space_after = Pt(1)
p_bl2.paragraph_format.keep_with_next = True
r_bl2 = p_bl2.add_run("BÀI LÀM (tiếp theo)")
r_bl2.bold = True
r_bl2.font.name = "Times New Roman"
r_bl2.font.size = Pt(11.5)

ExamDocxExporter._render_dot_lines(doc, count=27)

test_path = os.path.abspath('scratch/test_perfect_4p.docx')
doc.save(test_path)

word = win32com.client.Dispatch('Word.Application')
word.Visible = False
word.DisplayAlerts = 0
try:
    d = word.Documents.Open(test_path, ReadOnly=True)
    pgs = d.ComputeStatistics(2)
    print(f'>>> TOTAL COMPUTED PAGES FOR SINGLE EXAM: {pgs} <<<')
    for pg in range(1, pgs + 1):
        rng = d.GoTo(1, 1, pg)
        rng2 = d.GoTo(1, 1, pg + 1) if pg < pgs else d.Content
        rng.End = rng2.Start if pg < pgs else d.Content.End
        t = [l.strip() for l in rng.Text.split('\r') if l.strip()]
        first_l = t[0][:70] if t else '<empty>'
        print(f'  Page {pg} ({len(t)} lines): {first_l}')
    d.Close(False)
finally:
    word.Quit()
