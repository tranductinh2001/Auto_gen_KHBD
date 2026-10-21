import sys, os
sys.path.insert(0, os.getcwd())
sys.stdout.reconfigure(encoding='utf-8')

from core.exam_builder import ExamMatrixEngine, ExamDocxExporter
import docx

engine = ExamMatrixEngine(grade_level='Lớp 9')
dossier = engine.generate_exam_dossier()
out_file = 'scratch/test_date_and_align.docx'
ExamDocxExporter.export(dossier, out_file)
print('Export success!')

doc = docx.Document(out_file)

print('\n=== KIỂM TRA BẢN ĐẶC TẢ (Chữ ký & Ngày tháng) ===')
for i, t in enumerate(doc.tables):
    txt = ' '.join(c.text for row in t.rows for c in row.cells)
    if 'TỔ TRƯỞNG CHUYÊN MÔN' in txt:
        print(f'Table {i} (Bản đặc tả):')
        for r in t.rows:
            for ci, c in enumerate(r.cells):
                p = c.paragraphs[0]
                align = p.alignment
                print(f'  Col {ci} (align={align}): {p.text.splitlines()}')

print('\n=== KIỂM TRA HƯỚNG DẪN CHẤM (Chữ ký & Căn lề) ===')
for i, t in enumerate(doc.tables):
    txt = ' '.join(c.text for row in t.rows for c in row.cells)
    if 'Người ra đề, đáp án' in txt:
        print(f'Table {i} (Hướng dẫn chấm):')
        for r in t.rows:
            for ci, c in enumerate(r.cells):
                p = c.paragraphs[0]
                align = p.alignment
                print(f'  Col {ci} (align={align}): {p.text.splitlines()}')
